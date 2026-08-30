"""模块① 简历解析 Agent。

流程:PDF/DOCX/TXT → 文档解析 → LLM 结构化抽取(JSON)。

文档解析技术选型(对齐技术方案第八节):pdfminer.six(PDF)+ python-docx(DOCX),
轻量、纯 Python、Python 3.14 兼容。曾评估开源项目 DeepInterview 使用的 markitdown
(Apache-2.0),但其依赖 onnxruntime 在 Python 3.14 需源码编译,故未采用。
"""
import logging
from pathlib import Path
from typing import Any

from ..llm import chat_json

logger = logging.getLogger(__name__)


def file_to_text(path: str | Path) -> str:
    """将 PDF/DOCX/TXT/MD 转为纯文本(不上 LLM,可离线验证)。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_to_text(path)
    if suffix in (".docx", ".doc"):
        return _docx_to_text(path)
    # .txt / .md / 未知一律按纯文本读,编码兜底
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            return path.read_text(encoding=enc).strip()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError(f"无法读取文件编码: {path}")


def _pdf_to_text(path: Path) -> str:
    """pdfminer.six 提取 PDF 文本。"""
    from pdfminer.high_level import extract_text

    text = extract_text(str(path))
    text = (text or "").strip()
    if not text:
        raise ValueError("PDF 未提取到文本(可能是扫描件/图片型 PDF,MVP 暂不支持 OCR)")
    return text


def _docx_to_text(path: Path) -> str:
    """python-docx 提取 DOCX 文本(含表格)。"""
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("DOCX 未提取到文本")
    return text


SYSTEM_PROMPT = """你是简历解析专家。从候选人简历文本中抽取结构化信息。
要求:
1. 只抽取文本中真实存在的信息,不要编造;
2. 缺失的字段填 null 或空数组;
3. 严格输出 JSON,不要任何多余文字。

输出 JSON 结构(严格遵循):
{
  "basic": {"name": "", "age": "", "school": "", "major": ""},
  "education": [{"school": "", "degree": "", "major": "", "years": ""}],
  "experience": [{"company": "", "role": "", "duration": "", "summary": ""}],
  "projects": [{"name": "", "tech_stack": [], "achievements": ""}],
  "skills": {"programming": [], "tools": [], "soft": []}
}"""


def parse_resume_text(resume_text: str, max_chars: int = 8000) -> dict[str, Any]:
    """LLM 结构化抽取简历(内部函数,可单测)。"""
    return chat_json(
        system=SYSTEM_PROMPT,
        user=f"以下是候选人简历:\n\n{resume_text[:max_chars]}",
        temperature=0.1,
    )


def parse_resume_file(path: str | Path, max_chars: int = 8000) -> dict[str, Any]:
    """完整流水线:文件 → 文本 → 结构化 JSON。"""
    text = file_to_text(path)
    parsed = parse_resume_text(text, max_chars=max_chars)
    parsed["_source_text"] = text  # 原文留存,供面试官核验/评估引用
    return parsed
